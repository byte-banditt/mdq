"""Offline Excel export test."""

from openpyxl import load_workbook

from mdq.clean import rebuild_clean
from mdq.ingest import upsert_prices
from mdq.report import export_excel
from mdq.store import connect, init_db


def test_export_excel_sheets_and_issue_detail_count(tmp_path, prices_frame):
    db = str(tmp_path / "report.sqlite")
    init_db(db)
    run_id = "excel-test-run"
    with connect(db) as conn:
        upsert_prices(conn, prices_frame, "2024-01-03T00:00:00Z")
        conn.executemany(
            "INSERT INTO dq_issues(run_id,symbol,date,check_name,severity,detail,detected_at) "
            "VALUES(?,?,?,?,?,?,?)",
            [
                (run_id, "AAA.NS", "2024-01-01", "stale_price", "warning", "flat", "now"),
                (run_id, "BBB.NS", "2024-01-01", "return_outlier", "warning", "move", "now"),
            ],
        )
        rebuild_clean(conn, run_id)
        issue_count = conn.execute(
            "SELECT COUNT(*) FROM dq_issues WHERE run_id=?", (run_id,)
        ).fetchone()[0]

    path = export_excel(db, run_id, str(tmp_path / "excel"))
    workbook = load_workbook(path, read_only=True)

    assert workbook.sheetnames == ["Returns_Clean", "DQ_Summary", "DQ_Detail"]
    detail = workbook["DQ_Detail"]
    assert detail.max_row - 1 == issue_count
    workbook.close()
